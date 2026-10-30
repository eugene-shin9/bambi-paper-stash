#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
밤비의 서류뭉치 · PPT(16:9) 생성기  make_ppt.py
  python make_ppt.py --schema
  python make_ppt.py content.json --pal b3 --style base -o out.pptx
남이 열어 고칠 원본이므로 글꼴은 맑은 고딕(받는 PC에 있는 글꼴)으로 넣는다. 글꼴은 파일에 포함되지 않는다.
장 종류(slides[].type): cover 표지 · toc 목차 · divider 간지 · key 핵심 한 문장 · numbers 주요 수치 3칸 · compare 비교 2단 ·
  table 표 · bullets 개조식 · chart 막대 차트 · steps 진행 단계 · photo 사진·설명 · closing 마무리
"""
import sys, os, re, json, argparse
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)) if "__file__" in globals() else os.getcwd())
import bambi_common as B
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

FONT = "맑은 고딕"
SAMPLE = {
  "title": "2027 어린이 미술교육 운영 계획", "eyebrow": "교육팀 사업 보고 · 2026. 11. 12.", "dept": "교육팀 · 정다온", "doc": "미술교육 운영 계획",
  "slides": [
    {"type": "cover"},
    {"type": "toc", "items": [["현황", "지금 어디까지 와 있는지"], ["검토 내용", "무엇이 문제이고 왜 그런지"], ["향후 계획", "무엇을 언제 누가 할지"], ["행정 사항", "협조와 결정이 필요한 것"]]},
    {"type": "divider", "no": "01", "title": "현황", "sub": "이 장에서 다루는 것 · 추진 경과 · 현재 수치 · 남은 과제"},
    {"type": "key", "text": "가장 중요한 결론을 한 문장으로 적습니다", "mark": "한 문장", "points": ["근거 1 · 수치나 사실 한 줄", "근거 2 · 수치나 사실 한 줄", "근거 3 · 수치나 사실 한 줄"]},
    {"type": "numbers", "items": [["1,240", "명", "참여 인원", "전년 대비 12% 증가"], ["48", "회", "운영 횟수", "학기당 24회"], ["4.7", "점", "만족도", "5점 만점"]]},
    {"type": "compare", "left": {"title": "현행", "items": ["학기별 단발 프로그램", "강사 1인 운영"]}, "right": {"title": "개선", "items": ["학기 연계 4단계 과정", "강사 2인 + 보조"]}},
    {"type": "table", "title": "추진 일정", "head": ["단계", "내용", "담당", "기간"], "rows": [["1", "계획 수립", "교육팀", "11월"], ["2", "강사 섭외", "교육팀", "12월"]]},
    {"type": "bullets", "title": "개조식 본문", "items": [["항목 하나", ["세부 내용 가", "세부 내용 나"]], ["항목 둘", ["세부 내용 다"]]]},
    {"type": "chart", "title": "월별 참여 인원", "cats": ["3월", "4월", "5월", "6월"], "values": [120, 180, 210, 260], "unit": "명"},
    {"type": "steps", "items": [["계획", "11월"], ["섭외", "12월"], ["운영", "3~6월"], ["평가", "7월"]]},
    {"type": "photo", "title": "사진·설명", "image": "", "items": ["사진 설명 한 줄", "두 번째 설명"]},
    {"type": "closing", "summary": ["요약 1", "요약 2", "요약 3"], "next": "다음 단계 한 줄", "ask": "결정 요청 한 줄", "contact": "교육팀 053-000-0000"}
  ]
}
NOTE = """# content.json 규칙
# - slides 순서대로 장이 만들어진다. 필요한 장만 쓴다(표지·마무리는 권장). 같은 type을 여러 번 써도 된다.
# - 번호·글머리 기호는 쓰지 않는다. 원문에 없는 수치·날짜를 지어내지 않는다. image는 PNG/JPG 경로(없으면 "").
# - numbers는 3칸, steps는 3~5단계, chart 막대 8개까지, 표 열 5개까지.
"""
W, H = Cm(25.4), Cm(14.29); M = Cm(1.78); CW = W - 2*M

class Deck:
    def __init__(self, d, C, style):
        self.d, self.C, self.style = d, C, style; self.errs = []
        self.p = Presentation(); self.p.slide_width, self.p.slide_height = W, H; self.blank = self.p.slide_layouts[6]
        self.eb = B.head_on_dark(C)
    def c(self, k): return RGBColor.from_string(self.C[k])
    def slide(self, bg=None):
        s = self.p.slides.add_slide(self.blank)
        if bg: s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
        return s
    def box(self, s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
        sh = s.shapes.add_shape(shape, x, y, w, h)
        if fill: sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        else: sh.fill.background()
        if line: sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
        else: sh.line.fill.background()
        sh.shadow.inherit = False; return sh
    def text(self, s, x, y, w, h, lines, size=16, color="1C2430", bold=False, align="l", anchor="t", lh=1.2, space=4):
        """lines: str 또는 [(text, {size,color,bold})…] 줄 목록"""
        tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        if isinstance(lines, str): lines = [(lines, {})]
        for i, (t, o) in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[o.get("align", align)]
            para.line_spacing = lh; para.space_after = Pt(o.get("space", space))
            segs = t if isinstance(t, list) else [(t, {})]
            for st, so in segs:
                r = para.add_run(); r.text = st
                B.pptx_font(r, FONT, so.get("size", o.get("size", size)), so.get("bold", o.get("bold", bold)), so.get("color", o.get("color", color)))
        return tb
    def footer(self, s, n):
        self.text(s, M, Cm(13.0), Cm(12), Cm(0.6), f"{self.d.get('doc', self.d.get('title', ''))} · {self.d.get('dept', '').split(' · ')[0]}", 9, "5E6876")
        self.text(s, W - M - Cm(2), Cm(13.0), Cm(2), Cm(0.6), str(n), 9, "5E6876", align="r")
    def title(self, s, t):
        self.text(s, M, Cm(1.14), CW, Cm(1.2), t, 26, self.C["ink"], True)

    # ---- 장
    def cover(self, s, b):
        C = self.C; light = self.style == "light"; s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF" if light else C["main"])
        if self.style == "rich" and not light: self.box(s, W - Cm(8), -Cm(4.8), Cm(11.6), Cm(11.6), B.mix(C["sub"], C["main"], .4).__class__ and "%02X%02X%02X" % B.mix(C["sub"], C["main"], .4), shape=MSO_SHAPE.OVAL)
        fg = C["ink"] if light else "FFFFFF"; eb = C["main"] if light else self.eb
        self.text(s, M, Cm(3.2), CW, Cm(0.8), self.d.get("eyebrow", ""), 13, eb, True)
        self.text(s, M, Cm(4.3), CW, Cm(4), self.d.get("title", ""), 36, fg, True, lh=1.15)
        self.text(s, M, Cm(11.6), CW, Cm(0.8), self.d.get("dept", ""), 12, eb if not light else C["sub"])
    def toc(self, s, b):
        C = self.C; self.title(s, "목차"); items = b.get("items", [])[:6]; y = Cm(3.3); rh = Cm(1.75)
        for i, it in enumerate(items, 1):
            t, desc = (it + [""])[:2] if isinstance(it, list) else (str(it), "")
            self.text(s, M, y, Cm(1.6), rh, f"{i:02d}", 22, C["main"], True, anchor="m")
            self.text(s, M + Cm(1.8), y, CW - Cm(1.8), rh, [(t, {"size": 18, "bold": True, "space": 0}), (desc, {"size": 12, "color": C["sub"]})], anchor="m")
            self.box(s, M, y + rh, CW, Emu(1), line=C["line"], lw=0.5); y += rh + Cm(0.25)
        self.footer(s, self.n)
    def divider(self, s, b):
        C = self.C; s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string(C["tint"] if self.style != "light" else "FFFFFF")
        self.text(s, M, Cm(3.2), CW, Cm(2.5), b.get("no", ""), 54, C["main"], True)
        self.text(s, M, Cm(6.0), CW, Cm(1.6), b.get("title", ""), 32, C["ink"], True)
        self.text(s, M, Cm(8.0), CW, Cm(1), b.get("sub", ""), 14, C["sub"]); self.footer(s, self.n)
    def key(self, s, b):
        C = self.C; self.title(s, b.get("title", "핵심 메시지")); t = b.get("text", ""); mark = b.get("mark", "")
        self.box(s, M, Cm(3.3), CW, Cm(3.6), C["tint"] if self.style != "light" else None, None if self.style != "light" else C["line"])
        segs = [(t, {})]
        if mark and mark in t: a, bb = t.split(mark, 1); segs = [(a, {}), (mark, {"color": C["main"]}), (bb, {})]
        self.text(s, M + Cm(1), Cm(3.3), CW - Cm(2), Cm(3.6), [(segs, {"size": 28, "bold": True})], anchor="m", lh=1.25)
        y = Cm(7.6)
        for p in b.get("points", [])[:4]:
            self.box(s, M, y + Cm(0.28), Cm(0.3), Cm(0.3), C["main"], shape=MSO_SHAPE.OVAL); self.text(s, M + Cm(0.7), y, CW - Cm(0.7), Cm(0.9), p, 14, C["ink"]); y += Cm(1.0)
        self.footer(s, self.n)
    def numbers(self, s, b):
        C = self.C; self.title(s, b.get("title", "주요 수치")); items = b.get("items", [])[:3]; n = len(items); gap = Cm(0.6); cw = (CW - gap*(n-1))/n
        for i, it in enumerate(items):
            num, unit, name, desc = (it + ["", "", ""])[:4]; x = M + i*(cw + gap)
            self.box(s, x, Cm(3.6), cw, Cm(6.2), C["tint"] if self.style != "light" else None, None if self.style != "light" else C["line"])
            self.text(s, x + Cm(0.6), Cm(4.2), cw - Cm(0.8), Cm(2.2), [([(num, {"size": 48, "bold": True, "color": C["main"]}), (unit, {"size": 18, "color": C["sub"]})], {})])
            self.text(s, x + Cm(0.8), Cm(7.0), cw - Cm(1.6), Cm(2.4), [(name, {"size": 15, "bold": True}), (desc, {"size": 12, "color": C["sub"]})])
        self.footer(s, self.n)
    def compare(self, s, b):
        C = self.C; self.title(s, b.get("title", "비교")); cw = (CW - Cm(1.2))/2
        for k, side in enumerate(("left", "right")):
            d = b.get(side, {}); x = M + k*(cw + Cm(1.2)); hc = C["main"] if k == 1 else C["sub"]
            self.text(s, x, Cm(3.3), cw, Cm(1), d.get("title", ""), 16, "FFFFFF" if k == 1 else C["ink"], True, anchor="m")
            if k == 1: self.box(s, x, Cm(3.3), cw, Cm(1), hc).text_frame.text = ""; self.text(s, x + Cm(0.4), Cm(3.3), cw, Cm(1), d.get("title", ""), 16, "FFFFFF", True, anchor="m")
            else: self.box(s, x, Cm(4.3), cw, Emu(1), line=C["main"], lw=1.2)
            y = Cm(4.7)
            for it in d.get("items", [])[:6]:
                self.text(s, x, y, cw, Cm(1.1), it, 14, C["ink"], anchor="m"); self.box(s, x, y + Cm(1.1), cw, Emu(1), line=C["line"], lw=0.5); y += Cm(1.25)
        self.footer(s, self.n)
    def table(self, s, b):
        C = self.C; self.title(s, b.get("title", "표")); head = [str(x) for x in b.get("head", [])]; rows = [[str(c) for c in r] for r in b.get("rows", [])]; n = len(head) or len(rows[0])
        from pptx.oxml.ns import qn
        from lxml import etree
        tbl = s.shapes.add_table(len(rows) + (1 if head else 0), n, M, Cm(3.3), CW, Cm(0.9)*(len(rows)+1)).table
        tblPr = tbl._tbl.tblPr; tblPr.set("bandRow", "0"); tblPr.set("firstRow", "1" if head else "0")
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is None: sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
        NUM = re.compile(r"^[\d,.\s~월일()\-–%원명]+$")
        numcol = [i > 0 and all(NUM.match(r[i]) for r in rows if i < len(r) and r[i].strip()) for i in range(n)]
        def cell(ci, ri, txt, bold, color, bottom):
            c = tbl.cell(ri, ci); c.fill.background(); c.margin_left = c.margin_right = Cm(0.2); c.margin_top = c.margin_bottom = Cm(0.15); c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.text = ""; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT if numcol[ci] else PP_ALIGN.LEFT; r = p.add_run(); r.text = txt; B.pptx_font(r, FONT, 13, bold, color)
            tcPr = c._tc.get_or_add_tcPr()
            for tag in ("a:lnL", "a:lnR", "a:lnT"): el = etree.SubElement(tcPr, qn(tag), w="0"); etree.SubElement(el, qn("a:noFill"))
            ln = etree.SubElement(tcPr, qn("a:lnB"), w=str(int(12700*bottom[0]))); sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=bottom[1])
        ri = 0
        if head:
            for ci, h in enumerate(head): cell(ci, 0, h, True, C["main"], (1.2, C["main"]))
            ri = 1
        for r in rows:
            for ci in range(n): cell(ci, ri, r[ci] if ci < len(r) else "", False, C["ink"], (0.5, C["line"]))
            ri += 1
        self.footer(s, self.n)
    def bullets(self, s, b):
        C = self.C; self.title(s, b.get("title", "")); lines = []
        for it in b.get("items", [])[:8]:
            if isinstance(it, list): lines.append(("• " + it[0], {"size": 16, "bold": True, "space": 2})); lines += [("    – " + x, {"size": 14, "color": C["sub"], "space": 2}) for x in it[1][:4]]
            else: lines.append(("• " + str(it), {"size": 16, "space": 4}))
        self.text(s, M, Cm(3.3), CW, Cm(9), lines, lh=1.3); self.footer(s, self.n)
    def chart(self, s, b):
        C = self.C; self.title(s, b.get("title", "")); cd = CategoryChartData(); cd.categories = b.get("cats", []); cd.add_series(b.get("unit", ""), b.get("values", []))
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, M, Cm(3.0), CW, Cm(9.5), cd); ch = gf.chart; ch.has_legend = False; ch.has_title = False
        pl = ch.plots[0]; pl.gap_width = 80; pl.has_data_labels = True; dl = pl.data_labels; dl.position = XL_LABEL_POSITION.OUTSIDE_END; dl.font.size = Pt(12); dl.font.name = FONT
        ser = pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = self.c("main")
        va = ch.value_axis; va.visible = False; va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = self.c("line"); va.major_gridlines.format.line.width = Pt(0.5)
        ca = ch.category_axis; ca.tick_labels.font.size = Pt(12); ca.tick_labels.font.name = FONT; ca.format.line.color.rgb = self.c("line")
        self.footer(s, self.n)
    def steps(self, s, b):
        C = self.C; self.title(s, b.get("title", "진행 단계")); items = b.get("items", [])[:5]; n = len(items); cw = CW/n; r = Cm(0.75); cy = Cm(6.0)
        for i in range(n - 1): self.box(s, M + cw*i + cw/2 + r, cy + r - Emu(1), cw - 2*r, Emu(2), line=C["line"], lw=0.75)
        for i, it in enumerate(items):
            name, desc = (it + [""])[:2] if isinstance(it, list) else (str(it), ""); cx = M + cw*i + cw/2
            o = self.box(s, cx - r, cy, 2*r, 2*r, C["main"] if i == 0 else "FFFFFF", C["main"], MSO_SHAPE.OVAL, 1.5)
            tf = o.text_frame; tf.text = ""; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; rr = p.add_run(); rr.text = str(i+1); B.pptx_font(rr, FONT, 14, True, "FFFFFF" if i == 0 else C["main"]); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            self.text(s, M + cw*i, cy + 2*r + Cm(0.4), cw, Cm(2.4), [(name, {"size": 16, "bold": True, "align": "c", "space": 2}), (desc, {"size": 12, "color": C["sub"], "align": "c"})])
        self.footer(s, self.n)
    def photo(self, s, b):
        C = self.C; self.title(s, b.get("title", "")); pw, ph = Cm(10.4), Cm(8.6); img = b.get("image")
        if img and os.path.exists(img): s.shapes.add_picture(img, M, Cm(3.3), pw, ph)
        else: self.box(s, M, Cm(3.3), pw, ph, C["tint"] if self.style != "light" else None, None if self.style != "light" else C["line"])
        x = M + pw + Cm(1); lines = [(b.get("sub", ""), {"size": 20, "bold": True})] if b.get("sub") else []
        lines += [("• " + t, {"size": 15}) for t in b.get("items", [])[:6]]
        self.text(s, x, Cm(3.3), CW - pw - Cm(1), ph, lines, lh=1.3); self.footer(s, self.n)
    def closing(self, s, b):
        C = self.C; light = self.style == "light"; s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF" if light else C["main"])
        fg = C["ink"] if light else "FFFFFF"; eb = C["main"] if light else self.eb
        self.text(s, M, Cm(2.2), CW, Cm(1), b.get("title", "마무리"), 24, fg, True)
        self.text(s, M, Cm(3.8), CW, Cm(4), [("• " + t, {"size": 16}) for t in b.get("summary", [])[:4]], color=fg, lh=1.3)
        self.box(s, M, Cm(8.0), CW, Emu(1), line=("D5DDE7" if light else "FFFFFF"), lw=0.5)
        self.text(s, M, Cm(8.4), CW, Cm(2.6), [([("다음 단계  ", {"size": 13, "bold": True, "color": eb}), (b.get("next", ""), {"size": 14, "color": fg})], {}), ([("결정 요청  ", {"size": 13, "bold": True, "color": eb}), (b.get("ask", ""), {"size": 14, "color": fg})], {})], lh=1.3)
        self.text(s, M, Cm(11.6), CW, Cm(0.8), b.get("contact", ""), 11, eb if not light else C["sub"])

    def build(self, out):
        fn = dict(cover=self.cover, toc=self.toc, divider=self.divider, key=self.key, numbers=self.numbers, compare=self.compare, table=self.table, bullets=self.bullets, chart=self.chart, steps=self.steps, photo=self.photo, closing=self.closing)
        for i, b in enumerate(self.d.get("slides", []), 1):
            self.n = i; f = fn.get(b.get("type"))
            if not f: self.errs.append(f"{i}번째 장: 모르는 type {b.get('type')}"); continue
            f(self.slide(), b)
        self.p.save(out); return len(self.p.slides)

def validate(d):
    errs, warn = [], []
    if not str(d.get("title", "")).strip(): errs.append("title이 비어 있음")
    if not d.get("slides"): errs.append("slides가 비어 있음")
    for i, b in enumerate(d.get("slides", []), 1):
        t = b.get("type")
        if t == "numbers" and len(b.get("items", [])) != 3: errs.append(f"{i}번째 장(numbers)은 3칸이어야 함")
        if t == "chart" and (len(b.get("cats", [])) > 8 or len(b.get("cats", [])) != len(b.get("values", []))): errs.append(f"{i}번째 장(chart): 막대 8개 이하, cats와 values 개수 같게")
        if t == "table" and len(b.get("head", [])) > 5: errs.append(f"{i}번째 장(table): 열 5개까지")
        if t == "steps" and not 3 <= len(b.get("items", [])) <= 5: errs.append(f"{i}번째 장(steps): 3~5단계")
        for it in b.get("items", []) or []:
            if isinstance(it, str): B.strip_no(it, warn)
    if re.search(r"<[a-zA-Z/]", json.dumps(d, ensure_ascii=False)): errs.append("글 안에 <태그>가 있음")
    return errs, warn

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("content", nargs="?"); ap.add_argument("--schema", action="store_true"); ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--pal", default="b3"); ap.add_argument("--style", default="base", choices=["base", "light", "rich"]); ap.add_argument("-o", "--out")
    a = ap.parse_args()
    if a.schema: print(NOTE + json.dumps(SAMPLE, ensure_ascii=False, indent=2)); return
    if not a.content: ap.print_help(); return
    if a.pal not in B.PAL: B.log("색 id가 없음"); sys.exit(1)
    B.log("make_ppt.py sha256 " + B.self_hash(os.path.abspath(__file__)) + " (고치지 않았는지 대조용)")
    d = B.load_json(a.content); errs, warn = validate(d)
    for w in warn: B.log("[경고] " + w)
    if errs: B.fail(errs)
    if a.dry_run: B.log("검사 통과"); return
    out = a.out or os.path.join(B.OUTDIR, re.sub(r"[\\/:*?\"<>|\s]+", "_", d["title"])[:30] + f"_{B.PAL[a.pal]['name']}.pptx")
    dk = Deck(d, B.PAL[a.pal], a.style); n = dk.build(out)
    if dk.errs: B.fail(dk.errs)
    B.log(f"완료: {out} · {n}장 · 색 {B.PAL[a.pal]['name']} · 글꼴 맑은 고딕(받는 PC 글꼴, 파일에 포함 안 됨. Mac에서는 줄바꿈이 조금 달라질 수 있음)")
    print(out)

if __name__ == "__main__": main()
